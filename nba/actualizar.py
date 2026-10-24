"""Descarga los stats NBA y regenera el Excel y el dashboard.

Uso:  python3 nba/actualizar.py
Fuentes: basketball-reference.com (jugadores) y teamrankings.com (equipos).
"""
import datetime, json, os, subprocess, sys, time
import pandas as pd
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
import dvp
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

F="Arial"
def font(**k): return Font(name=F, **k)
thin=Side(style="thin",color="999999"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
C=Alignment(horizontal="center",vertical="center")
HDR=PatternFill("solid",fgColor="D9D9D9"); INP=PatternFill("solid",fgColor="FFFF00")
BLUE="0000FF"; GREEN="008000"

TEAM_NAMES=[("ATL","Atlanta Hawks",103.4,.27,.70,119.3,28.0,13.5),("BOS","Boston Celtics",97.0,.28,.74,107.2,24.0,13.0),
("BRK","Brooklyn Nets",97.5,.25,.72,112.2,26.5,13.3),("CHO","Charlotte Hornets",98.3,.29,.72,114.2,27.0,13.7),
("CHI","Chicago Bulls",103.6,.25,.72,119.4,27.5,14.0),("CLE","Cleveland Cavaliers",101.1,.27,.74,112.4,25.5,13.5),
("DAL","Dallas Mavericks",100.8,.27,.73,115.4,25.5,13.4),("DEN","Denver Nuggets",101.0,.30,.73,116.9,27.5,13.2),
("DET","Detroit Pistons",100.5,.30,.73,113.6,25.8,12.9),("GSW","Golden State Warriors",100.4,.30,.71,110.5,26.4,13.1),
("HOU","Houston Rockets",98.0,.35,.73,109.8,24.2,12.5),("IND","Indiana Pacers",100.0,.25,.72,115.1,26.7,12.9),
("LAC","LA Clippers",97.9,.26,.74,108.2,24.6,12.4),("LAL","Los Angeles Lakers",98.8,.25,.73,112.2,26.2,13.6),
("MEM","Memphis Grizzlies",104.9,.30,.73,116.9,26.6,13.8),("MIA","Miami Heat",97.3,.26,.74,110.0,25.4,13.5),
("MIL","Milwaukee Bucks",99.6,.26,.74,111.8,26.0,13.4),("MIN","Minnesota Timberwolves",97.4,.28,.73,109.3,24.9,12.8),
("NOP","New Orleans Pelicans",99.7,.29,.70,119.3,28.5,14.0),("NYK","New York Knicks",97.4,.29,.72,111.7,25.9,13.0),
("OKC","Oklahoma City Thunder",100.9,.28,.72,107.6,24.4,12.9),("ORL","Orlando Magic",96.9,.29,.73,105.5,23.6,11.5),
("PHI","Philadelphia 76ers",98.5,.26,.71,115.8,27.1,13.6),("PHO","Phoenix Suns",98.7,.26,.73,116.6,27.0,13.8),
("POR","Portland Trail Blazers",99.0,.32,.71,113.9,26.4,13.0),("SAC","Sacramento Kings",99.6,.27,.73,115.3,26.6,13.6),
("SAS","San Antonio Spurs",100.7,.27,.73,116.7,27.4,13.8),("TOR","Toronto Raptors",100.3,.29,.71,115.2,26.5,13.2),
("UTA","Utah Jazz",100.6,.30,.69,121.2,29.5,14.4),("WAS","Washington Wizards",100.9,.27,.70,120.4,28.4,14.2)]

SEASON_END=2027  # temporada 2026-27 -> NBA_2027 en basketball-reference
HERE=os.path.dirname(os.path.abspath(__file__))
RAW=os.path.join(HERE,"data","raw")
OUT=os.path.join(HERE,"Modelo_Props_NBA.xlsx")
HTML_TPL=os.path.join(HERE,"plantilla_dashboard.html")
HTML_OUT=os.path.join(HERE,"proyecciones-nba.html")
JSON_OUT=os.path.join(HERE,"data","nba_data.json")
HOY=datetime.date.today()
FECHA=f"{HOY.day}-{['ene','feb','mar','abr','may','jun','jul','ago','sep','oct','nov','dic'][HOY.month-1]}-{HOY.year}"

TR_STATS=[("pace","possessions-per-game"),("orb","offensive-rebounding-pct"),("drb","defensive-rebounding-pct"),
          ("pts","opponent-points-per-game"),("ast","opponent-assists-per-game"),("tpm","opponent-three-pointers-made-per-game"),
          ("ppg","points-per-game")]

def fetch(url,path):
    for intento in range(4):
        r=subprocess.run(["curl","-sS","-f","-m","60","-A","Mozilla/5.0","-o",path,url])
        if r.returncode==0: return
        time.sleep(2**(intento+1))
    sys.exit(f"No se pudo descargar {url}")

os.makedirs(RAW,exist_ok=True)
for _,slug in TR_STATS:
    fetch(f"https://www.teamrankings.com/nba/stat/{slug}",os.path.join(RAW,slug+".html")); time.sleep(1)
fetch(f"https://www.basketball-reference.com/leagues/NBA_{SEASON_END}_per_game.html",os.path.join(RAW,"per_game.html"))
TRMAP={"Atlanta":"ATL","Boston":"BOS","Brooklyn":"BRK","Charlotte":"CHO","Chicago":"CHI","Cleveland":"CLE","Dallas":"DAL","Denver":"DEN","Detroit":"DET","Golden State":"GSW","Houston":"HOU","Indiana":"IND","LA Clippers":"LAC","LA Lakers":"LAL","Memphis":"MEM","Miami":"MIA","Milwaukee":"MIL","Minnesota":"MIN","New Orleans":"NOP","New York":"NYK","Okla City":"OKC","Orlando":"ORL","Philadelphia":"PHI","Phoenix":"PHO","Portland":"POR","Sacramento":"SAC","San Antonio":"SAS","Toronto":"TOR","Utah":"UTA","Washington":"WAS"}
def num(v):
    v=str(v).strip()
    return float(v.replace("%",""))/100 if v.endswith("%") else float(v)
stats={}
for key,f in TR_STATS:
    t=pd.read_html(os.path.join(RAW,f+".html"))[0]
    col=t.columns[2]
    stats[key]={TRMAP[r.Team]:num(r[col]) for _,r in t.iterrows()}
full={a:n for a,n,*_ in TEAM_NAMES}
T=[(a,full[a],stats["pace"][a],stats["orb"][a],stats["drb"][a],stats["pts"][a],stats["ast"][a],stats["tpm"][a],stats["ppg"][a]) for a in sorted(full)]
pg=pd.read_html(os.path.join(RAW,"per_game.html"))[0]
pg=pg[pg.Rk.astype(str)!="Rk"]
pg=pg[~pg.Team.astype(str).str.contains("TM")]
pg=pg[pg.Player.astype(str)!="League Average"]
pg=pg.drop_duplicates("Player",keep="last").drop(columns=["Awards"],errors="ignore")
pg=pg.sort_values("Player",key=lambda c:c.str.lower()).reset_index(drop=True)
PROWS=[]
for _,r in pg.iterrows():
    row=[]
    for c in pg.columns:
        v=r[c]
        if c in ("Player","Team","Pos"): row.append(None if pd.isna(v) else str(v))
        else:
            try: row.append(None if pd.isna(v) or str(v)=="" else float(v))
            except ValueError: row.append(None)
    PROWS.append(row)

# Defensa por posición desde los box scores
BOX_DIR=os.path.join(HERE,"data","boxscores")
print("Bajando box scores nuevos…")
print(f"  {dvp.update_boxscores(SEASON_END,BOX_DIR)} partidos nuevos")
POSMAP={r[1]:dvp.main_position(r[4]) for r in PROWS}
DVP=dvp.compute_dvp(BOX_DIR,POSMAP)
DVP_GAMES=len(os.listdir(BOX_DIR))

wb=Workbook()
R=wb.active; R.title="Resultado"
SJ=wb.create_sheet("Stats Jugadores"); SE=wb.create_sheet("Stats Equipos")
PA=wb.create_sheet("Parametros"); IN=wb.create_sheet("Instrucciones")

# ---- Stats Jugadores (same column order as Basketball-Reference per game table)
h=["Rk","Player","Age","Team","Pos","G","GS","MP","FG","FGA","FG%","3P","3PA","3P%","2P","2PA","2P%","eFG%","FT","FTA","FT%","ORB","DRB","TRB","AST","STL","BLK","TOV","PF","PTS"]
for j,x in enumerate(h,1):
    c=SJ.cell(1,j,x); c.font=font(bold=True); c.fill=HDR; c.alignment=C; c.border=B
for i,row in enumerate(PROWS,2):
    for j,x in enumerate(row,1):
        c=SJ.cell(i,j,x); c.font=font(color=BLUE); c.border=B
        if h[j-1].endswith("%"): c.number_format="0.000"
SJ.column_dimensions["B"].width=26
for col in "ACDEFGHIJKLMNOPQRSTUVWXYZ": SJ.column_dimensions[col].width=7
SJ.freeze_panes="C2"
SJ["B1"].comment=Comment("Fuente: basketball-reference.com, NBA 2026-27 Per Game Stats, descargado "+FECHA+". Jugadores traspasados: solo su equipo actual. Para actualizar, pega la tabla Per Game desde A1.","Modelo")

# ---- Stats Equipos
h2=["Abrev","Equipo","Pace","ORB%","DRB%","Pts permitidos","Ast permitidas","3PM permitidos","Pts anotados"]
for j,x in enumerate(h2,1):
    c=SE.cell(1,j,x); c.font=font(bold=True); c.fill=HDR; c.alignment=C; c.border=B
for i,t in enumerate(T,2):
    for j,x in enumerate(t,1):
        c=SE.cell(i,j,x); c.font=font(color=BLUE); c.border=B
        if j in (4,5): c.number_format="0.0%"
        if j in (3,6,7,8,9): c.number_format="0.0"
SE.column_dimensions["B"].width=26
for col in "CDEFGHI": SE.column_dimensions[col].width=16
SE.freeze_panes="C2"
SE["A33"]="Fuente sugerida: teamrankings.com (Possessions per Game, Offensive/Defensive Rebounding %, Opponent Points/Assists/3PM per Game). Valores: temporada 2026-27, descargados "+FECHA+" (muestra muy pequeña al inicio de temporada)."
SE["A33"].font=font(italic=True,size=9)

# ---- Parametros
PA.column_dimensions["A"].width=38; PA.column_dimensions["B"].width=14; PA.column_dimensions["C"].width=60
prm=[("Ventaja de local (±)",0.01,"0.0%","Supuesto: +1% local / -1% visitante"),
("Edge mínimo para apostar",0.05,"0.0%","Supuesto: prob. modelo − prob. implícita"),
("CV Puntos",0.25,"0.00","Desv. estándar = proyección × CV (supuesto)"),
("CV Rebotes",0.35,"0.00",""),("CV Asistencias",0.40,"0.00",""),("CV Triples",0.60,"0.00",""),
("CV PRA",0.20,"0.00",""),("CV Pts+Reb",0.23,"0.00",""),("CV Pts+Ast",0.23,"0.00",""),
("CV Reb+Ast",0.30,"0.00",""),("CV Robos",0.80,"0.00",""),("CV Tapones",0.80,"0.00",""),
("Desv. estándar total del partido",12,"0.0","Supuesto: puntos"),
("Usar ajuste por posición","Sí","@","Sí = usa la defensa del rival contra la posición del jugador (hoja Defensa x Posicion). No = usa la defensa general del equipo.")]
PA["A1"]="Parámetro"; PA["B1"]="Valor"; PA["C1"]="Nota"
for c in PA[1]: c.font=font(bold=True); c.fill=HDR; c.border=B
for i,(a,v,f,nt) in enumerate(prm,2):
    PA.cell(i,1,a).font=font(); c=PA.cell(i,2,v); c.font=font(color=BLUE); c.fill=INP; c.number_format=f
    PA.cell(i,3,nt).font=font(italic=True,size=9)
PA["A16"]="Promedios de liga (calculados)"; PA["A16"].font=font(bold=True)
lg=[("Pace promedio","=AVERAGE('Stats Equipos'!C2:C31)","0.0"),
("ORB% promedio","=AVERAGE('Stats Equipos'!D2:D31)","0.0%"),
("DRB% promedio","=AVERAGE('Stats Equipos'!E2:E31)","0.0%"),
("Pts permitidos por posesión","=SUMPRODUCT('Stats Equipos'!F2:F31/'Stats Equipos'!C2:C31)/30","0.000"),
("Ast permitidas por posesión","=SUMPRODUCT('Stats Equipos'!G2:G31/'Stats Equipos'!C2:C31)/30","0.000"),
("3PM permitidos por posesión","=SUMPRODUCT('Stats Equipos'!H2:H31/'Stats Equipos'!C2:C31)/30","0.000"),
("Pts anotados por posesión","=SUMPRODUCT('Stats Equipos'!I2:I31/'Stats Equipos'!C2:C31)/30","0.000")]
for i,(a,f,nf) in enumerate(lg,17):
    PA.cell(i,1,a).font=font(); c=PA.cell(i,2,f); c.font=font(); c.number_format=nf
# names for clarity
HOME,EDGE="Parametros!$B$2","Parametros!$B$3"
CV={k:f"Parametros!$B${r}" for k,r in zip(["PTS","REB","AST","3PM","PRA","PR","PA","RA","STL","BLK"],range(4,14))}
LPACE,LORB,LDRB,LPTS,LAST,L3PM,LOFF=[f"Parametros!$B${r}" for r in range(17,24)]

# ---- Resultado
for col,w in zip("ABCDEFGHIJKLMNOP",[2,14,14,14,14,14,14,14,14,14,14,14,14,14,14,14]): R.column_dimensions[col].width=w
R.sheet_view.showGridLines=False
def box(rng,val,fill,size=16,bold=True,color="000000"):
    R.merge_cells(rng); a=R[rng.split(":")[0]]; a.value=val
    a.font=font(bold=bold,size=size,color=color); a.fill=PatternFill("solid",fgColor=fill); a.alignment=C
    for row in R[rng]:
        for c in row: c.border=B; c.fill=PatternFill("solid",fgColor=fill)
box("B2:C3","Jugador:","FFF2CC"); box("D2:J3","LaMelo Ball","FFFF00",color=BLUE)
box("K2:L3","xMP:","FCE4D6"); box("M2:N3",34,"FFFF00",color=BLUE)
box("B4:C5","Rival:","FFD966"); box("D4:J5","Washington Wizards","FFFF00",color=BLUE)
box("K4:L5","Local:","F4B183"); box("M4:N5","Sí","FFFF00",color=BLUE)
R["M2"].number_format="0.0"

dv1=DataValidation(type="list",formula1="='Stats Jugadores'!$B$2:$B$1000",allow_blank=False); R.add_data_validation(dv1); dv1.add("D2")
dv2=DataValidation(type="list",formula1="='Stats Equipos'!$B$2:$B$31"); R.add_data_validation(dv2); dv2.add("D4")
dv3=DataValidation(type="list",formula1='"Sí,No"'); R.add_data_validation(dv3); dv3.add("M4")

SJr="'Stats Jugadores'!"; SEr="'Stats Equipos'!"
def pj(col): return f"INDEX({SJr}${col}:${col},$P$8)"
def te(col,row): return f"INDEX({SEr}${col}:${col},{row})"
# helper rows (hidden-ish column P)
R["P7"]="Fila jugador"; R["P8"]=f"=MATCH($D$2,{SJr}$B:$B,0)"
R["P9"]="Fila equipo"; R["P10"]=f"=MATCH($F$8,{SEr}$A:$A,0)"
R["P11"]="Fila rival"; R["P12"]=f"=MATCH($D$4,{SEr}$B:$B,0)"
for a in ("P7","P9","P11"): R[a].font=font(size=8,color="808080")
R["P13"]="Posición"; R["P14"]=f'=IFERROR(LEFT({pj("E")},FIND("-",{pj("E")})-1),{pj("E")})'
for a in ("P8","P10","P12","P14"): R[a].font=font(size=8,color="808080")
R["P7"].font=font(size=8,color="808080"); R["P13"].font=font(size=8,color="808080")
DXP="'Defensa x Posicion'!"
USEPOS='Parametros!$B$15="Sí"'
def fx(stat): return f'IFERROR(INDEX({DXP}$C:$AF,$P$12,MATCH($P$14&" {stat}",{DXP}$C$1:$AF$1,0)),1)'

def table(row,headers,formulas,fmts):
    for j,(hd,f,nf) in enumerate(zip(headers,formulas,fmts)):
        c=R.cell(row,3+j,hd); c.font=font(bold=True,size=9); c.fill=HDR; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=B
        c=R.cell(row+1,3+j,f); c.font=font(color=GREEN if "!" in f and "Parametros" not in f else "000000"); c.number_format=nf; c.alignment=C; c.border=B
R["B7"]="Ajustes"; R["B7"].font=font(bold=True)
# row 7/8: C..N
table(7,["Local adj.","Min. prom.","Minuto %","Equipo jugador","PACE adj.","Defensa adj.","Off Reb adj.","Def Reb adj.","Ast adj.","3PM adj.","Factor base"],
 [f'=IF($M$4="Sí",1+{HOME},1-{HOME})', f"={pj('H')}", "=IFERROR($M$2/D8,0)", f"={pj('D')}",
  f"=(({te('C','$P$10')}+{te('C','$P$12')})/2)/{te('C','$P$10')}",
  f"=IF({USEPOS},{fx('PTS')},({te('F','$P$12')}/{te('C','$P$12')})/{LPTS})",
  f"=IF({USEPOS},{fx('REB')},(1-{te('E','$P$12')})/(1-{LDRB}))", f"=IF({USEPOS},{fx('REB')},(1-{te('D','$P$12')})/(1-{LORB}))",
  f"=IF({USEPOS},{fx('AST')},({te('G','$P$12')}/{te('C','$P$12')})/{LAST})", f"=IF({USEPOS},{fx('3PM')},({te('H','$P$12')}/{te('C','$P$12')})/{L3PM})",
  "=E8*G8*C8"],
 ["0.0%","0.0","0.0%","@","0.00","0.00","0.00","0.00","0.00","0.00","0.000"])
R["B10"]="Tiro"; R["B10"].font=font(bold=True)
table(10,["2PA proy.","2P%","3PA proy.","3P%","FTA proy.","FT%","ORB proy.","DRB proy.","Ast proy.","3PM proy.","Puntos adj."],
 [f"={pj('P')}*$M$8", f"=IFERROR({pj('O')}/{pj('P')},0)", f"={pj('M')}*$M$8", f"=IFERROR({pj('L')}/{pj('M')},0)",
  f"={pj('T')}*$M$8", f"=IFERROR({pj('S')}/{pj('T')},0)", f"={pj('V')}*$M$8*I8", f"={pj('W')}*$M$8*J8",
  f"={pj('Y')}*$M$8*K8", "=E11*F11*L8", "=(C11*D11*2+E11*F11*3+G11*H11)*H8"],
 ["0.0","0.0%","0.0","0.0%","0.0","0.0%","0.0","0.0","0.0","0.0","0.0"])

# Props table
R["B13"]="Props"; R["B13"].font=font(bold=True,size=12)
ph=["Mercado","Proyección","Promedio temporada","Línea","Cuota Over (US)","Cuota Under (US)","Prob. Over","Prob. Under","Prob. implícita Over","Prob. implícita Under","Edge Over","Edge Under","Pick"]
for j,hd in enumerate(ph,2):
    c=R.cell(14,j,hd); c.font=font(bold=True,size=9,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78")
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=B
R.row_dimensions[14].height=30
mk=[("Puntos","=M11",f"={pj('AD')}","PTS",26.5,-115,-105),
("Rebotes","=I11+J11",f"={pj('X')}","REB",5.5,-110,-110),
("Asistencias","=K11",f"={pj('Y')}","AST",7.5,-120,100),
("Triples","=L11",f"={pj('L')}","3PM",4.5,110,-130),
("PRA","=C15+C16+C17",f"=D15+D16+D17","PRA",39.5,-110,-110),
("Pts+Reb","=C15+C16","=D15+D16","PR",None,None,None),
("Pts+Ast","=C15+C17","=D15+D17","PA",None,None,None),
("Reb+Ast","=C16+C17","=D16+D17","RA",None,None,None),
("Robos",f"={pj('Z')}*$M$8*IF({USEPOS},{fx('STL')},1)",f"={pj('Z')}","STL",None,None,None),
("Tapones",f"={pj('AA')}*$M$8*IF({USEPOS},{fx('BLK')},1)",f"={pj('AA')}","BLK",None,None,None)]
for i,(m,pf,avg,k,ln,oo,ou) in enumerate(mk,15):
    r=i
    vals=[m,pf,avg,ln,oo,ou,
      f'=IF(E{r}="","",1-NORMDIST(E{r},C{r},MAX(0.5,C{r}*{CV[k]}),TRUE))',
      f'=IF(E{r}="","",1-H{r})',
      f'=IF(F{r}="","",IF(F{r}<0,-F{r}/(-F{r}+100),100/(F{r}+100)))',
      f'=IF(G{r}="","",IF(G{r}<0,-G{r}/(-G{r}+100),100/(G{r}+100)))',
      f'=IF(OR(H{r}="",J{r}=""),"",H{r}-J{r})',
      f'=IF(OR(I{r}="",K{r}=""),"",I{r}-K{r})',
      f'=IF(E{r}="","",IF(AND(L{r}<>"",L{r}>={EDGE}),"OVER",IF(AND(M{r}<>"",M{r}>={EDGE}),"UNDER","PASS")))']
    fm=["@","0.0","0.0","0.0","+0;-0","+0;-0","0.0%","0.0%","0.0%","0.0%","+0.0%;-0.0%","+0.0%;-0.0%","@"]
    for j,(v,f_) in enumerate(zip(vals,fm),2):
        c=R.cell(r,j,v); c.number_format=f_; c.border=B; c.alignment=C
        if j in (5,6,7): c.font=font(color=BLUE); c.fill=INP
        elif j==4: c.font=font(color=GREEN)
        else: c.font=font(bold=(j in (2,3,14)))
from openpyxl.formatting.rule import CellIsRule
R.conditional_formatting.add("N15:N24",CellIsRule(operator="equal",formula=['"OVER"'],fill=PatternFill("solid",fgColor="C6EFCE"),font=Font(name=F,bold=True,color="006100")))
R.conditional_formatting.add("N15:N24",CellIsRule(operator="equal",formula=['"UNDER"'],fill=PatternFill("solid",fgColor="FFC7CE"),font=Font(name=F,bold=True,color="9C0006")))

R["B26"]="Leyenda: celdas AMARILLAS con texto azul = entradas que tú llenas (jugador, rival, minutos, local, línea y cuotas). Verde = datos jalados de otras hojas. Negro = fórmulas."
R["B27"]="Datos reales 2026-27 (basketball-reference.com y teamrankings.com, "+FECHA+"). Las líneas y cuotas de ejemplo son ficticias: pon las de tu casa de apuestas."
for a in ("B26","B27"): R[a].font=font(italic=True,size=9)
R.freeze_panes="A6"

# Marcador proyectado del partido
R["B29"]="Marcador proyectado"; R["B29"].font=font(bold=True,size=12)
sh=["Posesiones","Pts equipo jugador","Pts rival","Total","Spread (equipo)","Línea total","Prob. Over total"]
for j,hd in enumerate(sh,2):
    c=R.cell(30,j,hd); c.font=font(bold=True,size=9,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78")
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=B
R.row_dimensions[30].height=30
sv=[f"=({te('C','$P$10')}+{te('C','$P$12')})/2",
    f"=B31*({te('I','$P$10')}/{te('C','$P$10')})*(({te('F','$P$12')}/{te('C','$P$12')})/{LOFF})*C8",
    f"=B31*({te('I','$P$12')}/{te('C','$P$12')})*(({te('F','$P$10')}/{te('C','$P$10')})/{LOFF})*(2-C8)",
    "=C31+D31","=D31-C31",225.5,
    f'=IF(G31="","",1-NORMDIST(G31,E31,Parametros!$B$14,TRUE))']
for j,(v,nf) in enumerate(zip(sv,["0.0","0.0","0.0","0.0","+0.0;-0.0","0.0","0.0%"]),2):
    c=R.cell(31,j,v); c.number_format=nf; c.border=B; c.alignment=C; c.font=font(bold=j in (3,4,5))
    if j==7: c.font=font(color=BLUE); c.fill=INP
R["B32"]="Pts = posesiones × pts anotados por posesión del equipo × (pts permitidos por posesión del rival ÷ promedio de liga) × local. Spread negativo = favorito."
R["B32"].font=font(italic=True,size=9)

# ---- Defensa x Posicion
DX=wb.create_sheet("Defensa x Posicion")
SLAB={"pts":"PTS","reb":"REB","ast":"AST","tpm":"3PM","stl":"STL","blk":"BLK"}
DX["A1"]="Abrev"; DX["B1"]="Equipo"
keys=[f"{p} {SLAB[s_]}" for p in dvp.POS for s_ in dvp.STATS]
for j,k in enumerate(keys):
    DX.cell(1,3+j,k); DX.cell(1,34+j,k+" /partido")
DX.cell(1,33,"Partidos")
for c in DX[1]:
    if c.value: c.font=font(bold=True,size=9); c.fill=HDR; c.alignment=Alignment(horizontal="center",wrap_text=True); c.border=B
for i,t in enumerate(T,2):
    DX.cell(i,1,t[0]).font=font(); DX.cell(i,2,t[1]).font=font()
    d=DVP.get(t[0],{})
    DX.cell(i,33,d.get("n",0)).font=font(color=BLUE)
    for pi,p in enumerate(dvp.POS):
        vals=d.get(p,[0]*6+[1]*6)
        for k in range(6):
            c=DX.cell(i,3+pi*6+k,vals[6+k]); c.number_format="0.00"; c.font=font(color=BLUE); c.border=B
            c=DX.cell(i,34+pi*6+k,vals[k]); c.number_format="0.0"; c.font=font(color=BLUE); c.border=B
from openpyxl.formatting.rule import ColorScaleRule
DX.conditional_formatting.add("C2:AF31",ColorScaleRule(start_type="num",start_value=0.8,start_color="F8696B",mid_type="num",mid_value=1,mid_color="FFFFFF",end_type="num",end_value=1.2,end_color="63BE7B"))
DX.column_dimensions["B"].width=24; DX.freeze_panes="C2"
DX.row_dimensions[1].height=30
DX["A33"]=("Factor = lo que el equipo permite a esa posición por posesión ÷ promedio de liga (1.20 = permite 20% más; verde = bueno para apostar Over). "
  f"Con pocos partidos se acerca a 1.00 (se mezcla con {dvp.SHRINK_GAMES} partidos de promedio de liga). Calculado con {DVP_GAMES} box scores de basketball-reference.com, {FECHA}. "
  "La posición de cada jugador es la de basketball-reference (la primera si tiene dos).")
DX["A33"].font=font(italic=True,size=9)
dvpar=DataValidation(type="list",formula1='"Sí,No"'); PA.add_data_validation(dvpar); dvpar.add("B15")

# ---- Instrucciones
lines=["CÓMO USAR EL MODELO DE PROPS NBA","",
"1. Actualiza 'Stats Jugadores': en basketball-reference.com > Season > Per Game Stats, copia la tabla y pégala en A1 (mismo orden de columnas: Rk, Player, Age, Team, Pos, G, GS, MP, FG ... PTS).",
"   Si un jugador fue traspasado aparece varias veces (2TM, 3TM): borra esas filas y deja la de su equipo actual.",
"• Defensa por posición: la hoja 'Defensa x Posicion' se calcula sola con nba/actualizar.py a partir de los box scores. Actívala o apágala en Parametros (Usar ajuste por posición).",
"2. Actualiza 'Stats Equipos' con teamrankings.com: Pace = Possessions per Game; ORB% = Offensive Rebounding %; DRB% = Defensive Rebounding %;",
"   Pts/Ast/3PM permitidos = Opponent Points / Opponent Assists / Opponent Three Pointers Made per Game. La columna Abrev debe coincidir con 'Team' de basketball-reference (CHO, PHO, BRK...).",
"3. En 'Resultado' elige jugador y rival (listas desplegables), escribe minutos esperados (xMP) y si juega de local (Sí/No).",
"4. Escribe la línea de la casa de apuestas y las cuotas americanas (ej. -115, +105). El modelo calcula probabilidad, edge y pick (OVER / UNDER / PASS).",
"","CÓMO CALCULA",
"• Minuto % = xMP / minutos promedio.   • PACE adj. = pace promedio de ambos equipos / pace del equipo del jugador.",
"• Defensa adj. = puntos permitidos por posesión del rival / promedio de liga (igual para asistencias y triples).",
"• Off Reb adj. = (1 − DRB% rival) / (1 − DRB% liga).   Def Reb adj. = (1 − ORB% rival) / (1 − ORB% liga).",
"• Factor base = Minuto % × PACE adj. × Local adj.  Cada estadística por partido se multiplica por el factor base y su ajuste.",
"• Puntos = (2PA×2P%×2 + 3PA×3P%×3 + FTA×FT%) × Defensa adj.",
"• Probabilidad Over = 1 − Normal(línea; proyección; proyección×CV). Los CV están en 'Parametros' y son supuestos editables.",
"• Prob. implícita: cuota negativa → −c/(−c+100); positiva → 100/(c+100). Edge = prob. modelo − prob. implícita.",
"","Herramienta educativa: no garantiza resultados. Apuesta con responsabilidad."]
IN.column_dimensions["A"].width=150
for i,l in enumerate(lines,1):
    c=IN.cell(i,1,l); c.font=font(bold=(i in (1,10)),size=14 if i==1 else 10)
wb.move_sheet("Instrucciones",offset=-3)
wb.active=0
wb.save(OUT)

# ---- Datos y dashboard
P_JSON=[]
for row in PROWS:
    d=dict(zip(h,row)); g=lambda k: round(float(d[k] or 0),2)
    P_JSON.append([d["Player"],d["Team"],d["Pos"] or "",int(d["G"] or 0),g("MP"),g("2P"),g("2PA"),g("3P"),g("3PA"),g("FT"),g("FTA"),g("ORB"),g("DRB"),g("AST"),g("STL"),g("BLK"),g("PTS")])
data={"fecha":FECHA,"p":P_JSON,"t":[list(t) for t in T],"dvp":DVP,"dvpGames":DVP_GAMES,"shrink":dvp.SHRINK_GAMES}
with open(JSON_OUT,"w") as f: json.dump(data,f,ensure_ascii=False,separators=(",",":"))
with open(HTML_TPL) as f: tpl=f.read()
with open(HTML_OUT,"w") as f: f.write(tpl.replace("__DATA__",json.dumps(data,ensure_ascii=False,separators=(",",":"))).replace("__FECHA__",FECHA).replace("__DVPGAMES__",str(DVP_GAMES)))
print(f"Listo ({FECHA}): {len(PROWS)} jugadores, {len(T)} equipos")
print(" ", OUT); print(" ", HTML_OUT)
